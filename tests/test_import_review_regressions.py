"""Regression tests for the CP3 (imports) independent review: one per finding."""

from __future__ import annotations

import io
import uuid
from datetime import datetime, timedelta, timezone
from decimal import Decimal

import pytest

from clawbot.domain.errors import ValidationFailed
from clawbot.domain.imports import files, values as v
from clawbot.persistence.db import transaction
from clawbot.workflows import import_worker
from tests.test_imports import (
    PRICE_LIST, SALES, _apply, _confirm_import, _csv, _mapping, _preview, _sales_upload, _upload, _xlsx, store,  # noqa: F401
)
from tests.test_suppliers import _approvable_supplier, _confirm


def _now_iso(**delta) -> str:
    return (datetime.now(timezone.utc) + timedelta(**delta)).isoformat()


# 1 — decompression bombs and oversized cells are refused before parsing
def test_xlsx_uncompressed_size_is_capped(monkeypatch):
    monkeypatch.setattr(files, "MAX_XLSX_UNCOMPRESSED_BYTES", 1_000)
    with pytest.raises(ValidationFailed, match="expands to more than"):
        files.parse_file("big.xlsx", _xlsx([["SKU", "Qty"], ["A", 1]]))


def test_oversized_cells_are_refused():
    with pytest.raises(ValidationFailed, match="longer than"):
        files.parse_file("a.csv", _csv([["SKU", "Qty"], ["A" * 3000, "1"]]))
    with pytest.raises(ValidationFailed, match="longer than"):
        files.parse_file("a.xlsx", _xlsx([["SKU", "Qty"], ["A" * 3000, 1]]))


# 2 — voids and cancellations are never demand
def test_unknown_sale_types_are_errors(store):
    rows = [["Date", "Transaction ID", "SKU", "Qty", "Event Type"],
            ["2026-10-01T10:00:00", "T-1", "FOOD-5", "1", "Void"],
            ["2026-10-01T10:00:00", "T-2", "FOOD-5", "1", "Cancelled"],
            ["2026-10-01T10:00:00", "T-3", "FOOD-5", "1", "Payment"],
            ["2026-10-01T10:00:00", "T-4", "FOOD-5", "1", "Refund"]]
    imp = _sales_upload(store, rows).json()
    p = _preview(store, imp)
    assert p["counts"] == {"rows": 4, "valid": 2, "with_errors": 2, "with_warnings": 0}
    assert "only sales and refunds" in p["rows"][0]["errors"][0]["message"]


# 3 — values the database would refuse are caught early; DB errors never strand an import
def test_overlong_values_are_row_errors(store):
    rows = [["Date", "Transaction ID", "SKU", "Qty"], ["2026-10-01T10:00:00", "T" * 300, "FOOD-5", "1"],
            ["2026-10-01T10:00:00", "T-2", "FOOD-5", "99999999999999"]]
    p = _preview(store, _sales_upload(store, rows).json())
    assert [r["errors"][0]["message"] for r in p["rows"]] == [
        "Transaction / order ID is longer than 255 characters", "Quantity '99999999999999' is too large"]


def test_database_error_on_one_row_does_not_sink_the_import(store, monkeypatch):
    from sqlalchemy.exc import DataError

    from clawbot.domain import sales

    real = sales.record_sales_line

    def flaky(session, **kw):
        if kw["external_order_id"] == "T-1002":
            raise DataError("INSERT", {}, Exception("value too long"))
        return real(session, **kw)

    monkeypatch.setattr(sales, "record_sales_line", flaky)
    strict = _apply(store, _sales_upload(store, SALES).json(), expect="failed")
    assert len(store.get("/v1/sales-lines").json()) == 0 and "nothing was applied" in strict["failure"]
    lenient = _apply(store, _sales_upload(store, SALES, name="again.csv").json(), accept_row_errors=True)
    assert lenient["summary"]["applied_counts"] == {"applied": 3, "error": 1}
    assert len(store.get("/v1/sales-lines").json()) == 3


def test_dead_lettered_job_marks_the_import_failed(store, monkeypatch):
    from clawbot.persistence.models import OutboxMessage

    imp = _sales_upload(store, SALES).json()
    assert _confirm_import(store, imp).status_code == 202
    with transaction() as s:
        s.query(OutboxMessage).update({"max_attempts": 1})
    monkeypatch.setattr(import_worker, "apply_import", lambda *a, **k: (_ for _ in ()).throw(RuntimeError("boom")))
    import_worker.run_once()
    after = store.get(f"/v1/imports/{imp['id']}").json()
    assert after["status"] == "failed" and "after 1 attempts" in after["failure"]


# 4 — re-upload respects context and status
def test_same_file_other_location_or_after_reject_is_a_new_import(store):
    other_loc = store.post("/v1/locations", {"name": "Second St", "timezone": "America/New_York"}).json()
    a = _sales_upload(store, SALES).json()
    b = _upload(store, "sales", _csv(SALES), location_id=other_loc["id"], source_system="square").json()
    assert a["id"] != b["id"] and b["context"]["location_id"] == other_loc["id"]
    store.post(f"/v1/imports/{a['id']}/reject", {"reason": "wrong store"})
    c = _sales_upload(store, SALES)
    assert c.status_code == 201 and c.json()["id"] not in (a["id"], b["id"])


# 5 — apply never does something nobody previewed
def test_change_between_confirm_and_apply_is_not_applied(store):
    supplier = _approvable_supplier(store)
    first = _upload(store, "supplier_price_list", _csv(PRICE_LIST), supplier_id=supplier["id"]).json()
    _mapping(store, first, price_basis="per_order_unit")
    _apply(store, first)
    item = next(i for i in store.get(f"/v1/suppliers/{supplier['id']}/items").json() if i["supplier_sku"] == "PH-LIT-20")

    again = [PRICE_LIST[0], ["PH-LIT-20", "Clumping litter 20lb", "036000291452", "12", "case", "$31.00", "2"]]
    imp = _upload(store, "supplier_price_list", _csv(again), supplier_id=supplier["id"]).json()
    _mapping(store, imp, price_basis="per_order_unit", update_pack_on_change=True)
    assert _confirm_import(store, imp).status_code == 202
    # A manager changes the pack before the worker runs.
    store.client.patch(f"/v1/supplier-items/{item['id']}", json={"version": item["version"], "units_per_order_uom": "24"},
                       headers=store.headers())
    import_worker.run_once()
    done = store.get(f"/v1/imports/{imp['id']}").json()
    # The person did not accept partial imports, so nothing is applied at all.
    assert done["status"] == "failed" and done["summary"]["changed_since_confirmation"] == 1
    assert store.get(f"/v1/supplier-items/{item['id']}").json()["units_per_order_uom"] == "24.0000"


def test_confirm_requires_restating_declared_settings(store):
    supplier = _approvable_supplier(store)
    imp = _upload(store, "supplier_price_list", _csv(PRICE_LIST), supplier_id=supplier["id"]).json()
    _mapping(store, imp, price_basis="per_order_unit")
    resp = _confirm_import(store, imp, declared={"price_basis": "per_base_unit"})
    assert resp.status_code == 422
    assert resp.json()["details"]["declared_settings_required"]["price_basis"] == "per_order_unit"


# 6 — unit labels are normalised
def test_cs_and_case_are_the_same_pack(store):
    supplier = _approvable_supplier(store)
    first = _upload(store, "supplier_price_list", _csv(PRICE_LIST), supplier_id=supplier["id"]).json()
    _mapping(store, first, price_basis="per_order_unit")
    _apply(store, first)
    relabelled = [PRICE_LIST[0], ["PH-LIT-20", "Clumping litter 20lb", "036000291452", "12", "CS", "$30.00", "2"]]
    imp = _upload(store, "supplier_price_list", _csv(relabelled), supplier_id=supplier["id"]).json()
    _mapping(store, imp, price_basis="per_order_unit")
    assert _preview(store, imp)["rows"][0]["errors"] == []
    assert _apply(store, imp)["summary"]["applied_counts"] == {"unchanged": 1}


# 7 — "$" is not only USD
@pytest.mark.parametrize("raw, currency, ok", [
    ("$30.00", "CAD", True), ("C$30", "CAD", True), ("US$30", "CAD", False), ("$30", "USD", True),
    ("¥3000", "CNY", True), ("€30", "USD", False), ("A$30", "AUD", True),
])
def test_currency_symbols(raw, currency, ok):
    if ok:
        assert v.decimal(raw, currency=currency) > 0
    else:
        with pytest.raises(v.CellError):
            v.decimal(raw, currency=currency)


def test_float_noise_is_refused_with_a_clear_message():
    with pytest.raises(v.CellError, match="more than 4 decimal places"):
        v.decimal("2.5833333333333335")


# 8 — line keys come from the file, so re-linking a barcode never double counts
def test_relinking_a_barcode_does_not_double_count(store):
    rows = [["Date", "Transaction ID", "UPC", "Qty"], ["2026-10-01T10:00:00", "T-1", "036000291452", "2"]]
    _apply(store, _sales_upload(store, rows).json())
    # The merchant moves the barcode to the food item and re-imports an overlapping export.
    for variant, barcode in ((store.litter, ""), (store.food, "036000291452")):
        current = store.get(f"/v1/variants/{variant['id']}").json()
        store.client.patch(f"/v1/variants/{variant['id']}", json={"version": current["version"], "barcode": barcode},
                           headers=store.headers())
    second = _sales_upload(store, rows + [["2026-10-02T10:00:00", "T-2", "036000291452", "1"]], name="b.csv").json()
    done = _apply(store, second, accept_row_errors=True)
    assert done["summary"]["applied_counts"] == {"error": 1, "applied": 1}  # T-1 flagged, not re-counted
    assert len([l for l in store.get("/v1/sales-lines").json() if l["external_order_id"] == "T-1"]) == 1


# 9 — repeated partial refunds are each recorded
def test_second_partial_refund_in_later_export_is_recorded(store):
    base = [["Date", "Transaction ID", "SKU", "Qty"], ["2026-10-01T10:00:00", "T-1", "FOOD-5", "3"],
            ["2026-10-02T10:00:00", "T-1", "FOOD-5", "-1"]]
    _apply(store, _sales_upload(store, base).json())
    later = base + [["2026-10-03T10:00:00", "T-1", "FOOD-5", "-1"]]
    assert _apply(store, _sales_upload(store, later, name="later.csv").json())["summary"]["applied_counts"] == {
        "unchanged": 2, "applied": 1}
    refunds = [l for l in store.get("/v1/sales-lines").json() if l["line_type"] == "refund"]
    assert len(refunds) == 2 and {l["sold_at"][:10] for l in refunds} == {"2026-10-02", "2026-10-03"}


# 10 — a newer count applies even if an older one was entered later
def test_counts_ordered_by_when_they_were_taken(store):
    a = _upload(store, "inventory", _csv([["SKU", "Qty"], ["FOOD-5", "10"]]), location_id=store.location_id).json()
    _mapping(store, a, as_of=_now_iso(hours=-3))
    _apply(store, a)
    b = _upload(store, "inventory", _csv([["SKU", "Qty"], ["FOOD-5", "8"]]), location_id=store.location_id).json()
    _mapping(store, b, as_of=_now_iso(hours=-2))
    assert _apply(store, b)["summary"]["applied_counts"] == {"applied": 1}
    levels = {p["variant_id"]: p["on_hand"] for p in store.get("/v1/inventory").json()}
    assert levels[store.food["id"]] == "8.0000"


# 11 — future-dated counts and sales are refused
def test_future_dates_refused(store):
    inv = _upload(store, "inventory", _csv([["SKU", "Qty"], ["FOOD-5", "5"]]), location_id=store.location_id).json()
    _mapping(store, inv, as_of=_now_iso(days=365))
    assert "in the future" in _preview(store, inv)["rows"][0]["errors"][0]["message"]
    rows = [["Date", "Transaction ID", "SKU", "Qty"], [_now_iso(days=2), "T-F", "FOOD-5", "1"]]
    assert "in the future" in _preview(store, _sales_upload(store, rows).json())["rows"][0]["errors"][0]["message"]


# 12 — uncalculated formulas are errors, not blanks
def test_uncalculated_formula_is_an_error():
    with pytest.raises(ValidationFailed, match="formula with no calculated value"):
        files.parse_file("f.xlsx", _xlsx([["SKU", "Qty"], ["A", "=10+2"]]))


# 13 — malformed CSV is a 4xx with a reason
def test_nul_bytes_and_binary_as_csv_refused():
    with pytest.raises(ValidationFailed, match="not a text CSV"):
        files.parse_file("a.csv", b"SKU,Qty\nA\x00,1\n")
    with pytest.raises(ValidationFailed, match="not a text CSV"):
        files.parse_file("a.csv", _xlsx([["SKU"], ["A"]]))


# 14 — malformed IDs in context are 422, not 500
def test_malformed_context_ids(store):
    resp = _upload(store, "inventory", _csv([["SKU", "Qty"], ["A", "1"]]), location_id="not-a-uuid")
    assert resp.status_code == 422


# §6.3 — one origin, one name, one key scheme
def test_free_text_origins_refused(store):
    resp = _upload(store, "sales", _csv(SALES), location_id=store.location_id, source_system="square-export")
    assert resp.status_code == 422


def test_order_already_recorded_by_another_path_is_not_counted_twice(store):
    from clawbot.domain import sales
    from clawbot.policies.permissions import system

    with transaction() as s:  # e.g. the Square connector recorded T-1001 with its own line UIDs
        sales.record_sales_line(
            s, tenant_id=uuid.UUID(store.tenant_id), actor=system("square-sync"),
            variant_id=uuid.UUID(store.litter["id"]), location_id=uuid.UUID(store.location_id),
            source_system="square", external_order_id="T-1001", external_line_key="LINE-UID-1",
            key_scheme="provider_line", sold_at=datetime(2026, 10, 1, 14, 15, tzinfo=timezone.utc),
            quantity=Decimal(2), line_type="sale",
        )
    done = _apply(store, _sales_upload(store, SALES).json(), accept_row_errors=True)
    assert done["summary"]["applied_counts"] == {"error": 2, "applied": 2}  # both T-1001 rows flagged
    errors = store.client.get(f"/v1/imports/{done['id']}/errors.csv", headers=store.headers()).text
    assert "through another path (provider_line)" in errors



# ---------------------------------------------------------------------------
# Second CP3 review round
# ---------------------------------------------------------------------------


def _xlsx_with_stale_dimension() -> bytes:
    import zipfile

    raw = _xlsx([["SKU", "Qty"], ["A", 1], ["B", 2]])
    src, out = zipfile.ZipFile(io.BytesIO(raw)), io.BytesIO()
    with zipfile.ZipFile(out, "w") as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                data = data.replace(b'<dimension ref="A1:B3"/>', b'<dimension ref="A1:B2"/>')
                assert b'ref="A1:B2"' in data
            dst.writestr(item, data)
    return out.getvalue()


def test_r2_stale_xlsx_dimension_keeps_every_row():
    parsed = files.parse_file("stale.xlsx", _xlsx_with_stale_dimension())
    assert [cells[0] for _, cells in parsed.rows] == ["A", "B"]


def test_r2_uncached_array_formula_is_an_error():
    from openpyxl import Workbook
    from openpyxl.worksheet.formula import ArrayFormula

    wb = Workbook()
    ws = wb.active
    ws.append(["SKU", "Qty"])
    ws["A2"] = "A"
    ws["B2"] = ArrayFormula("B2", "=SUM(1,2)")
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(ValidationFailed, match="formula with no calculated value"):
        files.parse_file("arr.xlsx", buf.getvalue())


def test_r2_cny_prefix_and_header_cap():
    assert v.decimal("CN¥300", currency="CNY") == 300
    with pytest.raises(ValidationFailed, match="column name is longer"):
        files.parse_file("h.csv", _csv([["S" * 3000], ["x"]]))


def test_r2_blank_line_id_is_a_preview_error(store):
    rows = [["Date", "Transaction ID", "Line ID", "SKU", "Qty"],
            ["2026-10-01T10:00:00", "T-1", "L1", "FOOD-5", "1"], ["2026-10-01T10:00:00", "T-1", "", "FOOD-5", "1"]]
    imp = _sales_upload(store, rows).json()
    assert _preview(store, imp)["counts"]["with_errors"] == 1


def test_r2_new_error_after_confirm_fails_whole_import_unless_accepted(store):
    from clawbot.domain import sales
    from clawbot.policies.permissions import system

    def record_t1002_elsewhere():
        with transaction() as s:
            sales.record_sales_line(
                s, tenant_id=uuid.UUID(store.tenant_id), actor=system("square-sync"),
                variant_id=uuid.UUID(store.litter["id"]), location_id=uuid.UUID(store.location_id),
                source_system="square", external_order_id="T-1002", external_line_key="UID-9",
                key_scheme="provider_line", sold_at=datetime(2026, 10, 2, 19, 40, tzinfo=timezone.utc),
                quantity=Decimal(1), line_type="sale",
            )

    imp = _sales_upload(store, SALES).json()
    assert _confirm_import(store, imp).status_code == 202
    record_t1002_elsewhere()
    import_worker.run_once()
    assert store.get(f"/v1/imports/{imp['id']}").json()["status"] == "failed"
    assert len(store.get("/v1/sales-lines").json()) == 1  # only the line recorded elsewhere


def test_r2_declared_settings_compare_by_meaning(store):
    inv = _upload(store, "inventory", _csv([["SKU", "Qty"], ["FOOD-5", "5"]]), location_id=store.location_id).json()
    when = datetime.now(timezone.utc) - timedelta(hours=1)
    _mapping(store, inv, as_of=when.isoformat(), timezone="America/Chicago")
    same_instant = when.astimezone(timezone(timedelta(hours=-5))).isoformat()
    ok = _confirm_import(store, inv, declared={"as_of": same_instant, "timezone": "America/Chicago"})
    assert ok.status_code == 202, ok.text  # omitted false/null defaults are fine
    sales_imp = _sales_upload(store, SALES).json()
    assert _confirm_import(store, sales_imp, declared={}).status_code == 202
    other = _sales_upload(store, SALES[:2], name="x.csv").json()
    _mapping(store, other, timezone="Asia/Seoul")
    assert _confirm_import(store, other, declared={}).status_code == 422  # an agent-set timezone must be restated


def test_r2_ordinals_count_rows_with_errors_too(store):
    rows = [["Date", "Transaction ID", "SKU", "Qty"],
            ["2026-10-01T10:00:00", "T-1", "NEW-SKU", "1"], ["2026-10-01T10:00:00", "T-1", "NEW-SKU", "2"]]
    imp = _sales_upload(store, rows).json()
    assert _confirm_import(store, imp, accept_row_errors=True).status_code == 422  # no valid rows at all
    store.post("/v1/variants", {"sku": "NEW-SKU", "name": "New thing"})
    keys = [r["interpretation"] for r in _preview(store, imp)["rows"]]
    assert len(keys) == 2
    _apply(store, imp)
    assert sorted(l["external_line_key"] for l in store.get("/v1/sales-lines").json()) == [
        "sku:NEW-SKU:sale:1", "sku:NEW-SKU:sale:2"]


def test_r2_crashed_worker_job_marks_import_failed(store):
    from clawbot.domain import outbox
    from clawbot.persistence.models import OutboxMessage

    imp = _sales_upload(store, SALES).json()
    assert _confirm_import(store, imp).status_code == 202
    with transaction() as s:
        s.query(OutboxMessage).update({"max_attempts": 1})
    with transaction() as s:  # worker claims, then the process dies
        outbox.claim_batch(s, lease=timedelta(seconds=-1), topics=("import.apply",))
    import_worker.run_once()  # lease expiry → dead; sweep marks the import
    assert store.get(f"/v1/imports/{imp['id']}").json()["status"] == "failed"


def test_upload_keeps_working_after_psycopg_prepares_the_statement(store):
    """ON CONFLICT on a partial index must use a literal predicate (bound params stop matching once prepared)."""
    for n in range(8):
        rows = [["Date", "Transaction ID", "SKU", "Qty"], ["2026-10-01T10:00:00", f"T-{n}", "FOOD-5", "1"]]
        assert _sales_upload(store, rows, name=f"f{n}.csv").status_code == 201
