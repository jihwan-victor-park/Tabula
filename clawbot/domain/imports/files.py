"""
File parsing: CSV and XLSX into (header, rows of strings).

XLSX is read with openpyxl in read-only mode using cached cell values; no
macro or formula is ever executed. A formula cell with no cached value (the
file was never calculated) is an error rather than an empty string.

Spec 16.6 file safety: the uncompressed size of a workbook is checked before
parsing (a 300 KB .xlsx can inflate to gigabytes), every cell is length-capped,
and the file's leading bytes must match its extension.
"""

from __future__ import annotations

import csv
import hashlib
import io
from dataclasses import dataclass

from clawbot.domain.errors import ValidationFailed
from clawbot.domain.imports.values import excel_cell_to_text

MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_XLSX_UNCOMPRESSED_BYTES = 60 * 1024 * 1024
MAX_ROWS = 100_000
MAX_COLUMNS = 200
MAX_CELL_CHARS = 2_000
_ZIP_MAGIC = b"PK\x03\x04"


@dataclass(frozen=True)
class ParsedFile:
    file_format: str
    sha256: str
    header: list[str]
    rows: list[tuple[int, list[str]]]  # (row number as shown in a spreadsheet, cells)


def sha256(content: bytes) -> str:
    return hashlib.sha256(content).hexdigest()


def _clean_header(header: list[str]) -> list[str]:
    cleaned = [h.strip() for h in header]
    if any(len(h) > MAX_CELL_CHARS for h in cleaned):
        raise ValidationFailed(f"A column name is longer than {MAX_CELL_CHARS:,} characters")
    if not any(cleaned):
        raise ValidationFailed("The first row must contain column names")
    seen: set[str] = set()
    for name in cleaned:
        if not name:
            raise ValidationFailed("Every column needs a name in the first row")
        if name.lower() in seen:
            raise ValidationFailed(f"Column name appears twice: {name!r}")
        seen.add(name.lower())
    if len(cleaned) > MAX_COLUMNS:
        raise ValidationFailed(f"Too many columns (max {MAX_COLUMNS})")
    return cleaned


def _rows(raw_rows, width: int) -> list[tuple[int, list[str]]]:
    rows: list[tuple[int, list[str]]] = []
    for number, cells in raw_rows:
        cells = [c if c is not None else "" for c in cells]
        if any(len(c) > MAX_CELL_CHARS for c in cells):
            raise ValidationFailed(f"Row {number} has a value longer than {MAX_CELL_CHARS:,} characters")
        if not any(c.strip() for c in cells):
            continue  # blank line
        if len(cells) > width and any(c.strip() for c in cells[width:]):
            raise ValidationFailed(f"Row {number} has more values than there are column names")
        cells = (cells + [""] * width)[:width]
        rows.append((number, cells))
        if len(rows) > MAX_ROWS:
            raise ValidationFailed(f"Too many rows (max {MAX_ROWS:,})")
    if not rows:
        raise ValidationFailed("The file has a header but no data rows")
    return rows


def _decode(content: bytes) -> str:
    if content.startswith(_ZIP_MAGIC) or b"\x00" in content:
        raise ValidationFailed("This is not a text CSV file (is it a workbook saved with a .csv name?)")
    try:
        return content.decode("utf-8-sig")
    except UnicodeDecodeError:
        raise ValidationFailed("The file is not UTF-8 text; re-export it as 'CSV UTF-8'")


def parse_csv(content: bytes) -> ParsedFile:
    text = _decode(content)
    try:
        dialect = csv.Sniffer().sniff(text[:20_000], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    # (csv.field_size_limit is process-global, so cell length is enforced in _rows instead.)
    reader = csv.reader(io.StringIO(text), dialect)
    try:
        header = next(reader)
        header = _clean_header(header)
        rows = _rows(((reader.line_num, row) for row in reader), len(header))
    except StopIteration:
        raise ValidationFailed("The file is empty")
    except csv.Error as exc:
        if "field larger than field limit" in str(exc):
            raise ValidationFailed(f"Row {reader.line_num} has a value longer than {MAX_CELL_CHARS:,} characters")
        raise ValidationFailed(f"The CSV cannot be read near line {reader.line_num}: {exc}")
    return ParsedFile("csv", sha256(content), header, rows)


def _check_zip(content: bytes) -> None:
    import zipfile

    if not content.startswith(_ZIP_MAGIC):
        raise ValidationFailed("The file is not a readable .xlsx workbook")
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            total = sum(info.file_size for info in archive.infolist())
    except zipfile.BadZipFile:
        raise ValidationFailed("The file is not a readable .xlsx workbook")
    if total > MAX_XLSX_UNCOMPRESSED_BYTES:
        raise ValidationFailed("The workbook expands to more than 60 MB; export the sheet as CSV instead")


def parse_xlsx(content: bytes, sheet: str | None = None) -> ParsedFile:
    import zipfile

    from openpyxl import load_workbook
    from openpyxl.worksheet.formula import ArrayFormula

    _check_zip(content)
    try:
        # Cached values for the data; a second, formula-preserving pass detects
        # formulas that were never calculated (their cached value is missing).
        values_book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        formula_book = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
    except (zipfile.BadZipFile, KeyError, ValueError, OSError):
        raise ValidationFailed("The file is not a readable .xlsx workbook")
    try:
        try:
            ws = values_book[sheet] if sheet else values_book.worksheets[0]
            fs = formula_book[ws.title]
        except (KeyError, IndexError):
            raise ValidationFailed(f"Sheet not found: {sheet}")
        # Read-only mode trusts the sheet's declared size, which some exporters
        # write stale; without this, rows past it vanish silently.
        ws.reset_dimensions()
        fs.reset_dimensions()
        pairs = zip(ws.iter_rows(values_only=True), fs.iter_rows(values_only=True))

        def text_row(number: int, values: tuple, formulas: tuple) -> list[str]:
            out = []
            for value, formula in zip(values, formulas):
                is_formula = isinstance(formula, ArrayFormula) or (isinstance(formula, str) and formula.startswith("="))
                if value is None and is_formula:
                    raise ValidationFailed(
                        f"Row {number} has a formula with no calculated value; open and save the file in Excel first"
                    )
                cell = excel_cell_to_text(value)
                if len(cell) > MAX_CELL_CHARS:
                    raise ValidationFailed(f"Row {number} has a value longer than {MAX_CELL_CHARS:,} characters")
                out.append(cell)
            return out

        try:
            first_values, first_formulas = next(pairs)
        except StopIteration:
            raise ValidationFailed("The sheet is empty")
        header = text_row(1, first_values, first_formulas)
        while header and not header[-1].strip():
            header.pop()  # trailing empty columns
        header = _clean_header(header)
        raw = ((i, text_row(i, vals, fms)) for i, (vals, fms) in enumerate(pairs, start=2))
        rows = _rows(raw, len(header))
    finally:
        values_book.close()
        formula_book.close()
    return ParsedFile("xlsx", sha256(content), header, rows)


def parse_file(file_name: str, content: bytes) -> ParsedFile:
    if len(content) > MAX_FILE_BYTES:
        raise ValidationFailed(f"File is larger than {MAX_FILE_BYTES // (1024 * 1024)} MB")
    if not content:
        raise ValidationFailed("The file is empty")
    lower = file_name.lower()
    if lower.endswith(".xlsx"):
        return parse_xlsx(content)
    if lower.endswith((".xls", ".xlsm", ".xlsb")):
        raise ValidationFailed("Save the workbook as .xlsx (macro-free) or CSV and upload that")
    if lower.endswith((".csv", ".tsv", ".txt")):
        return parse_csv(content)
    raise ValidationFailed("Upload a .csv or .xlsx file")
